"""
Parse Global Official Retail Price List (Retail) 1019.xlsx → Preview catalog data.
Does not write to Supabase / Zoho. Does not commit the source xlsx.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    r"c:\Users\Mark\Documents\xwechat_files\TCC---Mark_54d7\msg\file\2026-10\Global Official Retail Price List (Retail 공유용) 1019.xlsx"
)
OUT = ROOT / "data" / "imports" / "official-retail-1019"
IMG_DIR = OUT / "images"
THUMB_DIR = OUT / "thumbs"
REVIEW_DIR = OUT / "review"

COMING_SOON_TOKEN = "출시예정"


def slugify(text: str) -> str:
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")[:80] or "product"


def clean_whitespace(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def split_multilang_name(raw: str | None) -> dict:
    if not raw:
        return {"en": None, "kr": None, "cn": None, "jp": None, "raw": None, "aliasEn": None}
    text = clean_whitespace(str(raw))
    # Row 59 style: old name then -> new name
    alias_en = None
    if "->" in text or "→" in text:
        parts = re.split(r"(?:->|→)", text)
        # Prefer last English-looking Dermabell line
        candidates = [p.strip() for p in parts if p.strip()]
        en_candidates = []
        for part in candidates:
            for line in part.split("\n"):
                line = line.strip()
                if line.startswith("Dermabell") or re.match(r"^[A-Za-z]", line):
                    if "Dermabell" in line or re.search(r"[A-Za-z]{3,}", line):
                        en_candidates.append(line)
        if en_candidates:
            alias_en = en_candidates[0] if len(en_candidates) > 1 else None
            # final is last Dermabell
            derm = [e for e in en_candidates if e.startswith("Dermabell")]
            final_en = derm[-1] if derm else en_candidates[-1]
            # rebuild text focus on final block
            text_for_parse = parts[-1].strip()
        else:
            text_for_parse = text
            final_en = None
    else:
        text_for_parse = text
        final_en = None

    lines = [ln.strip() for ln in text_for_parse.split("\n") if ln.strip()]
    kr = cn = jp = en = None
    for line in lines:
        if line.startswith("->") or line.startswith("→"):
            line = re.sub(r"^(?:->|→)\s*", "", line)
        if re.search(r"[\uac00-\ud7af]", line) and not kr:
            kr = line
        elif re.search(r"[\u4e00-\u9fff]", line) and not cn:
            cn = line
        elif re.search(r"[\u3040-\u30ff]", line) and not jp:
            jp = line
        elif line.startswith("Dermabell") or (re.search(r"[A-Za-z]", line) and not en):
            if "Dermabell" in line or (en is None and re.match(r"^[A-Za-z]", line)):
                en = line if line.startswith("Dermabell") or en is None else en
                if line.startswith("Dermabell"):
                    en = line
    if final_en:
        if en and en != final_en:
            alias_en = alias_en or en
        en = final_en
    # strip leading arrows leftovers
    if en:
        en = re.sub(r"^(?:->|→)\s*", "", en).strip()
        en = re.sub(r"\s+", " ", en)
    return {"en": en, "kr": kr, "cn": cn, "jp": jp, "raw": text, "aliasEn": alias_en}


def parse_barcode_cell(value) -> dict:
    if value is None:
        return {"kind": "missing", "barcode": None, "components": None, "raw": None}
    if isinstance(value, (int, float)):
        barcode = str(int(value))
        return {"kind": "barcode", "barcode": barcode, "components": None, "raw": barcode}
    text = clean_whitespace(str(value))
    if COMING_SOON_TOKEN in text.replace(" ", ""):
        return {"kind": "coming_soon", "barcode": None, "components": None, "raw": text}
    if re.search(r"label", text, re.I):
        comps = {}
        for label, code in re.findall(r"(label|1st|2nd)\s*[:\n]?\s*(\d{8,14})", text, flags=re.I):
            comps[label.lower()] = code
        # also line-based
        if not comps:
            lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
            i = 0
            while i < len(lines):
                key = lines[i].lower()
                if key in ("label", "1st", "2nd") and i + 1 < len(lines) and re.fullmatch(r"\d{8,14}", lines[i + 1]):
                    comps[key] = lines[i + 1]
                    i += 2
                else:
                    i += 1
        return {
            "kind": "multi_component",
            "barcode": comps.get("label"),
            "components": comps,
            "raw": text,
        }
    m = re.search(r"(\d{8,14})", text)
    if m:
        return {"kind": "barcode", "barcode": m.group(1), "components": None, "raw": text}
    return {"kind": "unknown", "barcode": None, "components": None, "raw": text}


def pack_label_from_spec(spec: str | None) -> str | None:
    if not spec:
        return None
    text = clean_whitespace(spec)
    # Prefer metric line before fl.oz
    first = text.split("\n")[0].strip()
    first = re.sub(r"\s+", " ", first)
    return first


def classify_category(en_name: str | None, pack: str | None) -> str:
    hay = f"{en_name or ''} {pack or ''}".lower()
    if "peel" in hay and "wash-off" not in hay:
        return "Peels"
    if any(k in hay for k in ("mask", "modeling", "wash-off", "bio-gel")):
        return "Masks"
    if any(k in hay for k in ("cleanser", "cleansing foam", "softening solution", "milky soft")):
        return "Cleansers"
    if any(k in hay for k in ("toner", "pure water")):
        return "Toners"
    if any(k in hay for k in ("cream", "emulsion", "massage", "care gel", "aloe")):
        return "Creams & Emulsions"
    if any(k in hay for k in ("essence", "solution", "activator", "control solution")):
        return "Essences & Solutions"
    return "Other"


def product_group_key(en_name: str | None) -> str:
    if not en_name:
        return "unknown"
    # Remove trailing size tokens for grouping display name
    base = re.sub(
        r"\s*\d+(?:\.\d+)?\s*(?:ml|g|kg|oz|fl\.?\s*oz)\s*$",
        "",
        en_name,
        flags=re.I,
    ).strip()
    return base


def build_merged_name_map(ws) -> dict[int, str]:
    """Only inherit names within actual merge ranges for column B."""
    inherit: dict[int, str] = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col != 2 or rng.max_col != 2:
            continue
        top = ws.cell(rng.min_row, 2).value
        if not top:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            inherit[r] = str(top)
    return inherit


def extract_image_map(xlsx: Path) -> dict[int, str]:
    """Map Excel row number -> media filename using cell vm metadata (column C)."""
    with zipfile.ZipFile(xlsx) as zf:
        rels_root = ET.fromstring(zf.read("xl/richData/_rels/richValueRel.xml.rels"))
        rid_to_target = {
            rel.attrib["Id"]: rel.attrib["Target"].split("/")[-1]
            for rel in rels_root
            if rel.attrib.get("Id") and rel.attrib.get("Target")
        }
        rv_rel = ET.fromstring(zf.read("xl/richData/richValueRel.xml"))
        rel_ids = [
            el.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            for el in rv_rel
        ]
        # vm is 1-based index into futureMetadata / rich values
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        cells = re.findall(r'<c r="C(\d+)"[^>]*vm="(\d+)"', sheet)
        row_to_media: dict[int, str] = {}
        for row_s, vm_s in cells:
            row = int(row_s)
            vm = int(vm_s)  # 1-based
            idx = vm - 1
            if 0 <= idx < len(rel_ids):
                rid = rel_ids[idx]
                media = rid_to_target.get(rid)
                if media:
                    row_to_media[row] = media
        return row_to_media


def optimize_and_save(src_bytes: bytes, dest: Path, thumb: Path, max_edge: int = 1200) -> dict:
    dest.parent.mkdir(parents=True, exist_ok=True)
    thumb.parent.mkdir(parents=True, exist_ok=True)
    from io import BytesIO

    im = Image.open(BytesIO(src_bytes))
    im = im.convert("RGBA") if im.mode in ("P", "RGBA") else im.convert("RGB")
    w, h = im.size
    scale = min(1.0, max_edge / max(w, h))
    if scale < 1:
        im = im.resize((max(1, int(w * scale)), max(1, int(h * scale))), Image.Resampling.LANCZOS)
    im.save(dest, format="WEBP", quality=82, method=4)
    thumb_im = im.copy()
    thumb_im.thumbnail((240, 240), Image.Resampling.LANCZOS)
    thumb_im.save(thumb, format="WEBP", quality=75, method=4)
    return {"width": im.size[0], "height": im.size[1], "bytes": dest.stat().st_size}


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir(parents=True)
    IMG_DIR.mkdir()
    THUMB_DIR.mkdir()
    REVIEW_DIR.mkdir()

    wb = load_workbook(SOURCE, data_only=True)
    ws = wb[wb.sheetnames[0]]
    name_inherit = build_merged_name_map(ws)
    row_media = extract_image_map(SOURCE)

    with zipfile.ZipFile(SOURCE) as zf:
        media_files = {Path(n).name: zf.read(n) for n in zf.namelist() if n.startswith("xl/media/image")}

    variants = []
    image_map_rows = []
    issues = []

    for row in range(2, 68):
        a = ws.cell(row, 1).value
        b = ws.cell(row, 2).value
        if b is None and row in name_inherit:
            b = name_inherit[row]
            name_source = "merged_inherit"
        elif b is not None:
            name_source = "cell"
        else:
            name_source = "missing"
            issues.append({"row": row, "issue": "missing_name_outside_merge"})

        names = split_multilang_name(str(b) if b is not None else None)
        barcode_info = parse_barcode_cell(a)
        spec = ws.cell(row, 4).value
        pack = pack_label_from_spec(str(spec) if spec is not None else None)
        box_dim = ws.cell(row, 5).value
        unit_dim = ws.cell(row, 6).value
        unit_gw = ws.cell(row, 7).value
        box_gw = ws.cell(row, 8).value
        carton_qty = ws.cell(row, 9).value

        media_name = row_media.get(row)
        image_rel = None
        image_status = "missing"
        if media_name and media_name in media_files:
            stem = f"row-{row:02d}-{barcode_info.get('barcode') or 'nocode'}"
            stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", stem)
            dest = IMG_DIR / f"{stem}.webp"
            thumb = THUMB_DIR / f"{stem}.webp"
            meta = optimize_and_save(media_files[media_name], dest, thumb)
            image_rel = f"/imports/official-retail-1019/images/{dest.name}"
            image_status = "matched_cell_vm"
            image_map_rows.append(
                {
                    "row": row,
                    "excelMedia": media_name,
                    "vmMatched": True,
                    "local": image_rel,
                    "thumb": f"/imports/official-retail-1019/thumbs/{thumb.name}",
                    "bytes": meta["bytes"],
                    "width": meta["width"],
                    "height": meta["height"],
                    "barcode": barcode_info.get("barcode"),
                    "enName": names["en"],
                }
            )
        else:
            issues.append({"row": row, "issue": "image_unmatched", "media": media_name})
            image_map_rows.append(
                {
                    "row": row,
                    "excelMedia": media_name,
                    "vmMatched": False,
                    "local": None,
                    "barcode": barcode_info.get("barcode"),
                    "enName": names["en"],
                }
            )

        coming_soon = barcode_info["kind"] == "coming_soon"
        group_name = product_group_key(names["en"])
        variant_key = barcode_info.get("barcode") or f"row-{row}"
        if barcode_info["kind"] == "multi_component" and barcode_info.get("components"):
            variant_key = barcode_info["components"].get("label") or f"row-{row}"

        variant_id = "or1019-" + hashlib.sha1(f"{row}:{variant_key}".encode()).hexdigest()[:12]
        en = names["en"] or f"Untitled product (row {row})"
        category = classify_category(en, pack)

        variants.append(
            {
                "id": variant_id,
                "sourceRow": row,
                "barcode": barcode_info.get("barcode"),
                "barcodeKind": barcode_info["kind"],
                "barcodeComponents": barcode_info.get("components"),
                "barcodeRaw": barcode_info.get("raw"),
                "nameEn": en,
                "nameKr": names["kr"],
                "nameCn": names["cn"],
                "nameJp": names["jp"],
                "nameAliasEn": names["aliasEn"],
                "nameRaw": names["raw"],
                "nameSource": name_source,
                "groupNameEn": group_name,
                "packLabel": pack,
                "packSpecRaw": clean_whitespace(str(spec)) if spec is not None else None,
                "boxDimensionMm": clean_whitespace(str(box_dim)) if box_dim not in (None, "/") else None,
                "unitDimensionMm": clean_whitespace(str(unit_dim)) if unit_dim is not None else None,
                "unitGrossWeight": clean_whitespace(str(unit_gw)) if unit_gw is not None else None,
                "boxGrossWeight": clean_whitespace(str(box_gw)) if box_gw is not None else None,
                "cartonPackQty": int(carton_qty) if isinstance(carton_qty, (int, float)) else None,
                "category": category,
                "comingSoon": coming_soon,
                "purchasable": False,
                "priceCents": None,
                "images": [image_rel] if image_rel else [],
                "imageStatus": image_status,
                "slugHint": slugify(f"{group_name}-{pack or variant_key}"),
            }
        )

    # Group variants into products (same English group name)
    groups: dict[str, list] = defaultdict(list)
    for v in variants:
        groups[v["groupNameEn"]].append(v)

    products = []
    for group_name, items in groups.items():
        items = sorted(items, key=lambda x: x["sourceRow"])
        primary = items[0]
        product_id = "orp-" + hashlib.sha1(group_name.encode()).hexdigest()[:12]
        slug = slugify(group_name)
        # ensure unique slug
        existing_slugs = {p["slug"] for p in products}
        base_slug = slug
        n = 2
        while slug in existing_slugs:
            slug = f"{base_slug}-{n}"
            n += 1
        coming = all(i["comingSoon"] for i in items)
        products.append(
            {
                "id": product_id,
                "slug": slug,
                "name": group_name,
                "category": primary["category"],
                "comingSoon": coming,
                "variantIds": [i["id"] for i in items],
                "variants": items,
                "images": next((i["images"] for i in items if i["images"]), []),
                "source": "official-retail-1019-xlsx",
                "priceCents": None,
                "purchasable": False,
            }
        )

    products.sort(key=lambda p: (p["category"], p["name"]))

    # Special barcodes asserted
    melting = [v for v in variants if v["barcode"] in ("8809645782874", "8809645783581")]
    assert any(v["barcode"] == "8809645782874" and "70" in (v["packLabel"] or "") for v in melting), melting
    assert any(v["barcode"] == "8809645783581" and "220" in (v["packLabel"] or "") for v in melting), melting

    report = {
        "sourceFile": SOURCE.name,
        "sourceSheet": ws.title,
        "sourceRowRange": "2-67",
        "sourceRecordCount": len(variants),
        "productGroupCount": len(products),
        "comingSoonCount": sum(1 for v in variants if v["comingSoon"]),
        "imageMatchedCount": sum(1 for v in variants if v["images"]),
        "imageMissingCount": sum(1 for v in variants if not v["images"]),
        "multiComponentCount": sum(1 for v in variants if v["barcodeKind"] == "multi_component"),
        "mergedNameInheritCount": sum(1 for v in variants if v["nameSource"] == "merged_inherit"),
        "aliasRenameRows": [v["sourceRow"] for v in variants if v.get("nameAliasEn")],
        "issues": issues,
        "meltingMask": [
            {"barcode": v["barcode"], "pack": v["packLabel"], "row": v["sourceRow"]} for v in melting
        ],
        "dbStatus": "not_applied_supabase_readonly",
        "zohoStatus": "not_written",
        "previewLoader": "src/lib/catalog-official-retail.ts reads data/imports/official-retail-1019/catalog.json",
        "note": "No prices imported. New items are browse-only until priced in DB.",
    }

    catalog = {
        "version": 1,
        "generatedFrom": SOURCE.name,
        "products": products,
        "variants": variants,
    }

    (OUT / "catalog.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "image-map.json").write_text(json.dumps(image_map_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "import-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    # Markdown report
    md = [
        "# Official Retail 1019 import report",
        "",
        f"- Source: `{SOURCE.name}` / sheet `{ws.title}` rows 2–67",
        f"- Source records: **{report['sourceRecordCount']}**",
        f"- Product groups (front-end): **{report['productGroupCount']}**",
        f"- Coming soon: **{report['comingSoonCount']}**",
        f"- Images matched via cell vm: **{report['imageMatchedCount']}** / missing **{report['imageMissingCount']}**",
        f"- Multi-component barcode packs: **{report['multiComponentCount']}**",
        f"- Merged-name inherited rows: **{report['mergedNameInheritCount']}**",
        f"- Alias rename rows: {report['aliasRenameRows']}",
        "",
        "## Melting Mask barcodes",
        "",
    ]
    for m in report["meltingMask"]:
        md.append(f"- `{m['barcode']}` row {m['row']} · {m['pack']}")
    md += [
        "",
        "## Status",
        "",
        "- Supabase: **not applied** (read-only)",
        "- Zoho: **not written**",
        "- Prices: **none imported**",
        "- Preview loads `catalog.json` via official-retail catalog module",
        "",
    ]
    (OUT / "IMPORT_REPORT.md").write_text("\n".join(md), encoding="utf-8")

    # HTML review page for image alignment
    cards = []
    for row in image_map_rows:
        img = (
            f'<img src="../../..{row["thumb"]}" alt="" width="120" height="120" style="object-fit:contain;background:#f5f1ea" />'
            if row.get("thumb")
            else '<div style="width:120px;height:120px;background:#eee;display:flex;align-items:center;justify-content:center;font-size:11px">No image</div>'
        )
        cards.append(
            f"<tr><td>{row['row']}</td><td>{row.get('barcode') or ''}</td><td>{img}</td>"
            f"<td>{(row.get('enName') or '')}</td><td>{row.get('excelMedia') or ''}</td>"
            f"<td>{'OK' if row.get('vmMatched') else 'PENDING'}</td></tr>"
        )
    html = f"""<!doctype html><html><head><meta charset="utf-8"><title>Retail 1019 image review</title>
<style>body{{font-family:system-ui,sans-serif;padding:24px}} table{{border-collapse:collapse;width:100%}}
td,th{{border:1px solid #ddd;padding:8px;vertical-align:top;font-size:13px}} th{{background:#f7f4ef}}</style></head>
<body><h1>Official Retail 1019 — image alignment review</h1>
<p>Matched by Excel cell rich-value metadata (column C <code>vm</code>), not media file order.</p>
<table><thead><tr><th>Row</th><th>Barcode</th><th>Thumb</th><th>EN name</th><th>Excel media</th><th>Status</th></tr></thead>
<tbody>{''.join(cards)}</tbody></table></body></html>"""
    (REVIEW_DIR / "image-alignment.html").write_text(html, encoding="utf-8")

    print(json.dumps({k: report[k] for k in report if k != "issues"}, ensure_ascii=False, indent=2))
    print("issues", len(issues))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
