from openpyxl import load_workbook
from pathlib import Path
import sys

sys.stdout.reconfigure(encoding="utf-8")

path = Path(
    r"c:\Users\Mark\Documents\xwechat_files\TCC---Mark_54d7\msg\file\2026-10\Global Official Retail Price List (Retail 공유용) 1019.xlsx"
)
wb = load_workbook(path, data_only=True)
print("sheets", wb.sheetnames)
ws = wb[wb.sheetnames[0]]
print("active", ws.title, "dims", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
print("merged", len(ws.merged_cells.ranges))
for m in sorted(ws.merged_cells.ranges, key=lambda x: (x.min_row, x.min_col)):
    print(f" merge {m}")
print("--- header ---")
for col in range(1, 12):
    print(col, repr(ws.cell(1, col).value))
print("--- rows 2-25 ---")
for r in range(2, 26):
    vals = [ws.cell(r, c).value for c in range(1, 10)]
    print(r, vals)
print("images", len(getattr(ws, "_images", [])))
for i, img in enumerate(getattr(ws, "_images", [])):
    anchor = getattr(img, "anchor", None)
    fr = getattr(anchor, "_from", None)
    if fr is not None:
        print(f"img {i} from col={fr.col} row={fr.row} size={getattr(img, 'width', None)}x{getattr(img, 'height', None)}")
    else:
        print(f"img {i} anchor={anchor}")
print("--- rows 50-67 ---")
for r in range(50, 68):
    vals = [ws.cell(r, c).value for c in range(1, 10)]
    print(r, vals)
